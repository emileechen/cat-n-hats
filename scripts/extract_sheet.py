"""Extract the original Cats & Hats Google Sheet (exported as .xlsx) into clean data files.

Usage:
    python3 scripts/extract_sheet.py data/source/cats-and-hats-sheet-2026-10-08.xlsx data

Writes:
    data/catalog/hats.json          released hats (starter, basic, animals, april fools)
    data/catalog/cat_colours.json   all cat colours (paintbrushes)
    data/catalog/shop.json          shard shop listings (never went live)
    data/catalog/milestones.json    set completion rewards (from the Hat Brochure text)
    data/catalog/planned.json       unreleased sets, to-do list, player-written flavour text
    data/history/ledger.csv         the Records tab, with corrections and issue flags
    data/history/players.json       each closet as recorded vs. as derived from the ledger
    data/history/fashion_shows.json the two dress-up events that ran

Requires openpyxl.
"""

import csv
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

PLAYER_TABS = [
    "Aralie", "Ember", "Luv007", "Pinkshadow", "Ryuukokoro", "Saturnalia", "SeaCrest",
    "Silvanon", "Springacres", "Laviee", "Sunshine Warrior", "Wildfilly94", "XLostxDestinyX", "YLO",
]
REWARD_ONLY_HATS = {"AC19", "AC20"}
HAT_ID = re.compile(r"A[A-D]\d\d")
COLOUR_ID = re.compile(r"\*[A-C]\d\d")
IMAGE_REF = re.compile(r"IMAGE\((Cats|Hats)!\$?([A-Z])\$?(\d+)", re.I)
IMAGE_URL = re.compile(r'IMAGE\("(https://[^"]+)"', re.I)
ART_SET_DIRS = {"Starter": "starter", "Basic": "basic", "Animals": "animals", "April Fools 2022": "april-fools-2022"}


def load(path):
    return openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True)


def rows(ws, min_row=2):
    for row in ws.iter_rows(min_row=min_row, values_only=True):
        if any(v is not None for v in row):
            yield row


def reveal_parts(text):
    """'reaches into the box and pulls out......... a Red Ribbon!\\nhttps://...' -> 'Red Ribbon'"""
    if not text or "pulls out" not in text:
        return None
    m = re.search(r"pulls out\.+\s*(?:an?|some)\s+(.+?)!", text)
    return m.group(1) if m else None


def art_slug(url, suffix):
    """'https://i.postimg.cc/x/barette-pink-m.png' -> 'barrette-pink' (matches the file names in art/)."""
    name = url.rsplit("/", 1)[-1]
    return name[: -len(suffix)].replace("barette-", "barrette-") if name.endswith(suffix) else None


def extract_hats(wv):
    hats, planned = [], []
    for r in rows(wv["Hats"]):
        set_name, name, hat_id, desc, _owned, model_url, icon_url = r[:7]
        box_roll, reveal = r[9], r[10]
        if not hat_id:
            if set_name and name and name != "?":
                planned.append({"set": set_name, "name": name})
            continue
        if set_name == "Starter":
            obtain = "starter"
        elif box_roll is not None:
            obtain = "mystery_box"
        elif hat_id in REWARD_ONLY_HATS:
            obtain = "milestone_reward"
        else:
            obtain = "event"
        hats.append({
            "id": hat_id,
            "name": name,
            "display_name": reveal_parts(reveal),
            "set": set_name,
            "description": (desc or "").strip() or None,
            "obtain": obtain,
            "box_roll": int(box_roll) if box_roll is not None else None,
            "art": f"hats/{ART_SET_DIRS[set_name]}/{art_slug(model_url, '-m.png')}",
            "model_url": model_url,
            "icon_url": icon_url,
        })
    return hats, planned


def extract_colours(wv):
    obtain = {"Starter": "starter", "Basic": "basic_reward_or_shop", "Shop Exclusive": "shop_exclusive"}
    return [
        {
            "id": cid, "name": name, "set": set_name, "description": desc,
            "obtain": obtain.get(set_name), "art": f"cats/{art_slug(img, '.png')}",
            "image_url": img, "icon_url": icon,
        }
        for set_name, name, cid, desc, img, icon in (r[:6] for r in rows(wv["Cats"]))
        if cid
    ]


def extract_shop(wv):
    return [
        {"type": r[0].lower(), "item_id": r[3], "name": r[2], "cost_shards": int(r[8]),
         "available_flag": r[7] is True}
        for r in rows(wv["Shop"]) if r[3]
    ]


MILESTONES = [
    {"set": "Basic", "threshold": 6, "label": "50%", "rewards": [{"type": "colour_choice", "pool": "Basic"}, {"type": "shards", "amount": 50}],
     "source": "Hat Brochure: 'Collecting 6 Hats (50%) -- Choice of Basic paintbrush & 50 shards'", "awarded_in_history": 5},
    {"set": "Basic", "threshold": 12, "label": "100%", "rewards": [{"type": "undecided", "note": "brochure shows '?'"}, {"type": "shards", "amount": 200}],
     "source": "Hat Brochure: 'Collecting 12 Hats (100%) -- ? & 200 shards'", "awarded_in_history": 0},
    {"set": "Animals", "threshold": 9, "label": "50%", "rewards": [{"type": "hat", "id": "AC19"}, {"type": "shards", "amount": 50}],
     "source": "Hat Brochure: 'Collecting 9 Hats (50%) -- Exclusive hat & 50 shards'", "awarded_in_history": 6},
    {"set": "Animals", "threshold": 19, "label": "100%", "rewards": [{"type": "hat", "id": "AC20"}, {"type": "shards", "amount": 200}],
     "source": "Hat Brochure: 'Collecting 19 Hats (100%) -- Exclusive hat & 200 shards'",
     "awarded_in_history": 0,
     "issue": "Only 18 Animals hats are in the box pool. 19 only works if AC19 (the 50% reward) counts. Either that is intended or the threshold should be 18."},
]


def extract_planning(wf, wv):
    ws = wv["Planning"]
    columns = {}
    for col in range(1, 14):
        header = ws.cell(row=4, column=col).value
        if not header:
            continue
        items = [ws.cell(row=r, column=col).value for r in range(5, 30)]
        columns[header] = [i for i in items if i]
    flavour = []
    for cell in ("E2", "J2"):
        text = ws[cell].value or ""
        author, _, body = text.partition(":")
        flavour += [{"by": author.strip(), "text": t.strip()} for t in body.split(";") if t.strip()]
    return {
        "todo": [t for t in (ws["A2"].value or "").splitlines() if t.strip()],
        "planning_tab_sets": columns,
        "cat_flavour_text_suggestions": flavour,
    }


def image_ref_to_id(formula, hat_rows, colour_rows):
    m = IMAGE_REF.search(formula or "")
    if not m:
        return None
    sheet, _col, row = m.group(1).lower(), m.group(2), int(m.group(3))
    return (colour_rows if sheet == "cats" else hat_rows).get(row)


def extract_closets(wf, wv, hat_rows, colour_rows):
    players = {}
    for name in PLAYER_TABS:
        f, v = wf[name], wv[name]
        p = {"tab": name, "cat_name": None, "portrait_url": None, "current_outfit": None,
             "outfits": [], "recorded": {"shards": v["N3"].value, "colours": [], "hats": {}}}
        for row in f.iter_rows():
            for c in row:
                if not isinstance(c.value, str):
                    continue
                url = IMAGE_URL.search(c.value)
                if url and "-big." in url.group(1):
                    p["portrait_url"] = url.group(1)
                    below = v.cell(row=c.row + 1, column=c.column).value
                    p["cat_name"] = below if isinstance(below, str) else None
                    icons_row = c.row + 2 if p["cat_name"] else c.row + 1
                    p["current_outfit"] = {
                        "colour": image_ref_to_id(f.cell(row=icons_row, column=c.column).value, hat_rows, colour_rows),
                        "hat": image_ref_to_id(f.cell(row=icons_row, column=c.column + 3).value, hat_rows, colour_rows),
                    }
                elif url and re.search(r"-\d{3}\.png", url.group(1)):
                    date = v.cell(row=c.row + 1, column=c.column).value
                    p["outfits"].append({
                        "date": date.date().isoformat() if date else None,
                        "image_url": url.group(1),
                        "colour": image_ref_to_id(f.cell(row=c.row + 2, column=c.column).value, hat_rows, colour_rows),
                        "hat": image_ref_to_id(f.cell(row=c.row + 2, column=c.column + 1).value, hat_rows, colour_rows),
                    })
                val = v[c.coordinate].value
                if isinstance(val, str) and COLOUR_ID.fullmatch(val):
                    p["recorded"]["colours"].append(val)
                elif isinstance(val, str) and HAT_ID.fullmatch(val):
                    qty = v.cell(row=c.row, column=c.column + 1).value
                    if qty:
                        p["recorded"]["hats"][val] = int(qty)
        players[name] = p
    return players


def extract_ledger(wv, players):
    canonical = {name.lower(): name for name in players}
    out = []
    for i, r in enumerate(wv["Records"].iter_rows(min_row=2, max_col=8, values_only=True), start=2):
        date, user, source, item_in, item_out, s_in, s_out, note = r
        if not user:
            continue
        issues, player, item = [], canonical.get(user.lower(), user), item_in
        if player != user:
            issues.append(f"name typed as '{user}'; the closet's balance formula skipped this row")
        if item_in == "AB18":
            item = "AC18"
            issues.append("AB18 doesn't exist; closet shows AC18")
        if source == "Event" and item_in and item_in.startswith("AD") and item_in not in ("AD01", "AD02"):
            issues.append(f"{item_in} is a personalized Jester Hat (art/hats/april-fools-2022/jester-*); "
                          "the Hats tab and closets only had one Jester slot (AD02)")
        if source == "Mystery" and not item_in:
            issues.append("box opened but no hat recorded")
        out.append({
            "row": i, "date": date.date().isoformat(), "player": player, "source": source,
            "item_in": item, "item_out": item_out,
            "shards_in": int(s_in or 0), "shards_out": int(s_out or 0),
            "note": note, "item_in_as_written": item_in if item != item_in else None,
            "issues": "; ".join(issues) or None,
        })
    return out


def reconcile(players, ledger):
    hats, colours, shards = defaultdict(Counter), defaultdict(set), Counter()
    for e in ledger:
        p = e["player"]
        if e["item_in"]:
            if e["item_in"].startswith("*"):
                colours[p].add(e["item_in"])
            else:
                hats[p][e["item_in"]] += 1
        if e["item_out"]:
            hats[p][e["item_out"]] -= 1
        shards[p] += e["shards_in"] - e["shards_out"]
    for name, p in players.items():
        derived_hats = {k: n for k, n in sorted(hats[name].items()) if n}
        p["from_ledger"] = {"shards": shards[name], "colours": sorted(colours[name]), "hats": derived_hats}
        rec = p["recorded"]
        diffs = [
            f"{h}: closet {rec['hats'].get(h, 0)} vs ledger {derived_hats.get(h, 0)}"
            for h in sorted(set(rec["hats"]) | set(derived_hats))
            if rec["hats"].get(h, 0) != derived_hats.get(h, 0)
        ]
        if set(rec["colours"]) != colours[name]:
            diffs.append(f"colours: closet {sorted(rec['colours'])} vs ledger {sorted(colours[name])}")
        if rec["shards"] != shards[name]:
            diffs.append(f"shards: closet {rec['shards']} vs ledger {shards[name]}")
        p["discrepancies"] = diffs


def extract_fashion_shows(wv, ledger):
    ws = wv["Fashion Show"]
    shows = []
    for date_cell, text_cell in (("F16", "F17"), ("F3", "F4")):
        date = ws[date_cell].value.date().isoformat()
        theme, _, reward = (ws[text_cell].value or "").partition("\n")
        shows.append({
            "date": date,
            "theme": theme.removeprefix("Theme: ").strip(),
            "reward": reward.removeprefix("Reward: ").strip(),
            "entrants": [
                {"player": e["player"], "colour_chosen": e["item_in"], "shards": e["shards_in"]}
                for e in ledger if e["source"] == "Fashion Show" and e["date"] == date
            ],
        })
    return shows


def main(xlsx, out_dir):
    wf, wv = load(xlsx)
    out = Path(out_dir)
    (out / "catalog").mkdir(parents=True, exist_ok=True)
    (out / "history").mkdir(parents=True, exist_ok=True)

    hats, planned_from_hats = extract_hats(wv)
    colours = extract_colours(wv)
    hat_rows = {r: wv["Hats"].cell(row=r, column=3).value for r in range(2, 40)}
    colour_rows = {r: wv["Cats"].cell(row=r, column=3).value for r in range(2, 22)}
    planned = extract_planning(wf, wv)
    planned["hats_tab_unreleased"] = planned_from_hats

    players = extract_closets(wf, wv, hat_rows, colour_rows)
    ledger = extract_ledger(wv, players)
    reconcile(players, ledger)

    def dump(rel, data):
        (out / rel).write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n")

    dump("catalog/hats.json", hats)
    dump("catalog/cat_colours.json", colours)
    dump("catalog/shop.json", extract_shop(wv))
    dump("catalog/milestones.json", MILESTONES)
    dump("catalog/planned.json", planned)
    dump("history/players.json", players)
    dump("history/fashion_shows.json", extract_fashion_shows(wv, ledger))
    with open(out / "history/ledger.csv", "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(ledger[0]))
        w.writeheader()
        w.writerows(ledger)

    print(f"{len(hats)} hats, {len(colours)} colours, {len(ledger)} ledger rows, {len(players)} closets")
    for name, p in players.items():
        if p["discrepancies"]:
            print(f"  {name}: {'; '.join(p['discrepancies'])}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
